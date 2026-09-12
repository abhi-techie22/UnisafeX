#!/usr/bin/env python3
"""Fetch real public monument images from Wikimedia Commons into the workbook.

The script fills empty Image_URL_3..Image_URL_8 slots first. Any remaining empty
slots stay available for AI generation through the queue sheet.
"""

from __future__ import annotations

import argparse
import json
import time
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAIN_SHEET = "Monuments_Full_Profile"
QUEUE_SHEET = "Image_Generation_Queue"
COMMONS_API = "https://commons.wikimedia.org/w/api.php"
USER_AGENT = "UniSafeX/1.0 tourism-image-enrichment (educational dataset preparation)"

IMAGE_URL_COLUMNS = [f"Image_URL_{index}" for index in range(3, 9)]
SUPPORTED_MIME = {"image/jpeg", "image/png", "image/webp"}


def clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def header_map(sheet) -> dict[str, int]:
    return {
        clean(cell.value): cell.column
        for cell in sheet[1]
        if clean(cell.value)
    }


def ensure_columns(sheet, columns: list[str]) -> dict[str, int]:
    headers = header_map(sheet)
    for column in columns:
        if column not in headers:
            next_column = sheet.max_column + 1
            sheet.cell(row=1, column=next_column, value=column)
            headers[column] = next_column
    return headers


def commons_query(search_text: str, limit: int, timeout_seconds: int) -> dict[str, Any]:
    params = {
        "action": "query",
        "format": "json",
        "generator": "search",
        "gsrnamespace": "6",
        "gsrsearch": search_text,
        "gsrlimit": str(limit),
        "prop": "imageinfo",
        "iiprop": "url|size|mime|extmetadata",
    }
    url = f"{COMMONS_API}?{urllib.parse.urlencode(params)}"
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=timeout_seconds) as response:
        return json.loads(response.read().decode("utf-8"))


def metadata_value(metadata: dict[str, Any], key: str) -> str:
    item = metadata.get(key)
    if isinstance(item, dict):
        return clean(item.get("value"))
    return ""


def normalize_credit(value: str) -> str:
    # Commons metadata may include simple HTML. Keep credits compact and sheet-safe.
    text = value.replace("<span", " <span").replace("<", " <")
    output = []
    inside = False
    for char in text:
        if char == "<":
            inside = True
            continue
        if char == ">":
            inside = False
            continue
        if not inside:
            output.append(char)
    return " ".join("".join(output).split())[:240]


def image_candidates(row: dict[str, object], per_query_limit: int, timeout_seconds: int) -> list[dict[str, str]]:
    name = clean(row.get("Monument_Name"))
    locality = clean(row.get("Locality"))
    district = clean(row.get("District"))
    state = clean(row.get("State_UT"))
    category = clean(row.get("Category"))
    queries = [
        " ".join(part for part in [name, locality, state, "India monument"] if part),
        " ".join(part for part in [name, district, state, "India"] if part),
        " ".join(part for part in [name, category, "India"] if part),
    ]

    seen_urls: set[str] = set()
    results: list[dict[str, str]] = []
    for query in queries:
        if not query:
            continue
        try:
            payload = commons_query(query, per_query_limit, timeout_seconds)
        except Exception:
            continue
        pages = payload.get("query", {}).get("pages", {})
        if not isinstance(pages, dict):
            continue
        for page in pages.values():
            infos = page.get("imageinfo") or []
            if not infos:
                continue
            info = infos[0]
            url = clean(info.get("url"))
            mime = clean(info.get("mime"))
            width = int(info.get("width") or 0)
            height = int(info.get("height") or 0)
            if not url or url in seen_urls:
                continue
            if mime not in SUPPORTED_MIME or width < 400 or height < 300:
                continue
            metadata = info.get("extmetadata") or {}
            license_short = metadata_value(metadata, "LicenseShortName")
            usage_terms = metadata_value(metadata, "UsageTerms")
            artist = normalize_credit(metadata_value(metadata, "Artist"))
            source = metadata_value(metadata, "ObjectName") or clean(page.get("title"))
            seen_urls.add(url)
            results.append(
                {
                    "url": url,
                    "source": f"Wikimedia Commons: {source}",
                    "license": license_short or usage_terms or "Check Wikimedia license",
                    "credit": artist or "Wikimedia Commons contributor",
                }
            )
        if len(results) >= len(IMAGE_URL_COLUMNS):
            break
        time.sleep(0.02)
    return results[: len(IMAGE_URL_COLUMNS)]


def update_queue_for_fetched(workbook, filled_global_ids: set[str]) -> None:
    if QUEUE_SHEET not in workbook.sheetnames:
        return
    sheet = workbook[QUEUE_SHEET]
    headers = header_map(sheet)
    global_id_col = headers.get("Global_ID")
    status_col = headers.get("Status")
    source_mode_col = headers.get("Source_Mode")
    if not global_id_col or not status_col:
        return
    for row_index in range(2, sheet.max_row + 1):
        global_id = clean(sheet.cell(row=row_index, column=global_id_col).value)
        if global_id not in filled_global_ids:
            continue
        if source_mode_col:
            sheet.cell(row=row_index, column=source_mode_col, value="BOTH")
        if clean(sheet.cell(row=row_index, column=status_col).value) == "Pending":
            sheet.cell(row=row_index, column=status_col, value="Pending only if image slot remains empty")


def fetch_into_workbook(
    input_path: Path,
    output_path: Path,
    max_rows: int | None,
    per_query_limit: int,
    timeout_seconds: int,
) -> None:
    workbook = load_workbook(input_path)
    if MAIN_SHEET not in workbook.sheetnames:
        raise SystemExit(f"Missing sheet: {MAIN_SHEET}")
    sheet = workbook[MAIN_SHEET]
    extra_columns = []
    for index in range(3, 9):
        extra_columns.extend(
            [
                f"Image_Source_{index}",
                f"Image_License_{index}",
                f"Image_Credit_{index}",
                f"Image_Verification_Status_{index}",
            ]
        )
    headers = ensure_columns(sheet, extra_columns)
    base_headers = header_map(sheet)

    existing_image_cols = [
        base_headers[column]
        for column in [f"Image_URL_{index}" for index in range(1, 9)]
        if column in base_headers
    ]
    filled_global_ids: set[str] = set()
    processed = 0

    for row_index in range(2, sheet.max_row + 1):
        if max_rows is not None and processed >= max_rows:
            break
        row = {
            name: sheet.cell(row=row_index, column=col).value
            for name, col in base_headers.items()
        }
        empty_slots = [
            column
            for column in IMAGE_URL_COLUMNS
            if column in base_headers and not clean(sheet.cell(row=row_index, column=base_headers[column]).value)
        ]
        if not empty_slots:
            continue
        existing_urls = {
            clean(sheet.cell(row=row_index, column=column).value)
            for column in existing_image_cols
            if clean(sheet.cell(row=row_index, column=column).value)
        }
        candidates = [
            item
            for item in image_candidates(row, per_query_limit, timeout_seconds)
            if item["url"] not in existing_urls
        ]
        if candidates:
            filled_global_ids.add(clean(row.get("Global_ID")))
        for slot, candidate in zip(empty_slots, candidates):
            image_index = slot.rsplit("_", 1)[-1]
            sheet.cell(row=row_index, column=base_headers[slot], value=candidate["url"])
            sheet.cell(row=row_index, column=headers[f"Image_Source_{image_index}"], value=candidate["source"])
            sheet.cell(row=row_index, column=headers[f"Image_License_{image_index}"], value=candidate["license"])
            sheet.cell(row=row_index, column=headers[f"Image_Credit_{image_index}"], value=candidate["credit"])
            sheet.cell(row=row_index, column=headers[f"Image_Verification_Status_{image_index}"], value="Fetched - needs admin review")
        processed += 1
        if processed % 50 == 0:
            print(f"Processed {processed} rows; fetched images for {len(filled_global_ids)} monuments", flush=True)
        time.sleep(0.02)

    update_queue_for_fetched(workbook, filled_global_ids)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Saved: {output_path}")
    print(f"Rows processed: {processed}")
    print(f"Monuments with fetched images: {len(filled_global_ids)}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--max-rows", type=int, default=None, help="Optional limit for testing/resume batches.")
    parser.add_argument("--per-query-limit", type=int, default=10)
    parser.add_argument("--timeout-seconds", type=int, default=8)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    fetch_into_workbook(
        args.input.expanduser(),
        args.output.expanduser(),
        args.max_rows,
        args.per_query_limit,
        args.timeout_seconds,
    )


if __name__ == "__main__":
    main()
